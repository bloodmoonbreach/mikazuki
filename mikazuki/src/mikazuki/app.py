"""mikazuki web app. Data lives in $MIKAZUKI_DATA (default ./data)."""

from __future__ import annotations

import datetime as dt
import mimetypes
import os
import re
from pathlib import Path

import yaml
from fastapi import FastAPI, Form, Request
from fastapi.responses import FileResponse, HTMLResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates

import cvtracks
from cvtracks import i18n as cv_i18n
from cvtracks.cli import EXAMPLES, _file_stem
from cvtracks.loader import load_master, load_track, load_tracks
from cvtracks.model import ModelError
from cvtracks.render import RENDERERS, Options, render
from cvtracks.select import resolve
from remote_eligibility import LLMConfig, Profile, assess

from . import __version__, apply, feed, keys, places, security
from .net import FetchError, fetch
from .store import Store, StoreError, valid_track_id

HERE = Path(__file__).parent
mimetypes.add_type("font/woff2", ".woff2")  # slim images lack it; fonts would be served as octet-stream
store = Store(os.environ.get("MIKAZUKI_DATA", "data"))
applications = apply.AppStore(store.root)
app = FastAPI(title="mikazuki", version=__version__, docs_url=None, redoc_url=None, openapi_url=None)
app.mount("/static", StaticFiles(directory=HERE / "static"), name="static")
templates = Jinja2Templates(directory=HERE / "templates")


def _safe_url(url) -> str:
    """Only http(s) links reach an href; anything else (javascript:, data:, …) becomes a dead link."""
    u = str(url or "").strip()
    return u if u.lower().startswith(("http://", "https://")) else "#"


templates.env.filters["safe_url"] = _safe_url
auth = security.install(app, store.root)


@app.on_event("startup")
def _start_scheduler() -> None:
    import os as _os
    places.use_data_folder(store)
    keys.apply(store)
    if _os.environ.get("MIKAZUKI_SCHEDULER", "1") != "0":
        feed.start_scheduler(store)
        places.ensure(store)

ICONS = {"yes": "✅", "likely": "🟢", "unclear": "❔", "unlikely": "🟠", "no": "❌"}
TRACK_TEMPLATE = """\
id: {id}
title: {{en: "{title}", es: "{title}"}}
languages: [en, es]
summary: null            # a key from `summaries` in master.yaml
include_tags: []         # empty = everything; e.g. [soc]
exclude_tags: [unpublished]
sections: [summary, experience, projects, certifications, skills, education, languages]
limits: {{pages: 1, bullets: 4}}
keywords: []             # terms screeners look for; lint warns if the CV never mentions them
"""


def page(request: Request, name: str, **ctx) -> HTMLResponse:
    ctx.setdefault("flash", request.query_params.get("msg"))
    ctx.setdefault("error", request.query_params.get("err"))
    return templates.TemplateResponse(request, name, {"csrf": request.state.csrf, "nav": name.split(".")[0],
                                                      "version": __version__, "auth": auth.configured, **ctx})


def back(url: str, msg: str | None = None, err: str | None = None) -> RedirectResponse:
    from urllib.parse import urlencode
    q = urlencode({k: v for k, v in (("msg", msg), ("err", err)) if v})
    return RedirectResponse(url + (("&" if "?" in url else "?") + q if q else ""), status_code=303)


def load():
    if not store.has_master:
        raise StoreError("no master.yaml yet")
    master = load_master(store.path("master.yaml"))
    tracks = load_tracks(store.path("tracks"))
    return master, tracks


def llm_key() -> str | None:
    """The API key saved by the setup (data/.llm-key, readable only by the app), if any."""
    f = store.root / ".llm-key"
    try:
        return f.read_text(encoding="utf-8").strip() or None
    except OSError:
        return None


def profile_and_llm() -> tuple[Profile, LLMConfig | None, dict]:
    prof = store.profile()
    profile = Profile.parse(prof.get("countries") or ["ES"], prof.get("citizenships") or [], prof.get("languages"))
    llm_cfg = prof.get("llm") or {}
    cfg = None if llm_cfg.get("provider") == "none" else LLMConfig.from_env(llm_cfg)
    if cfg and not cfg.api_key and llm_cfg.get("api_key_file"):
        cfg.api_key = llm_key()
    return profile, cfg, prof


# ---------------------------------------------------------------- pages


@app.get("/login", response_class=HTMLResponse)
def login_form(request: Request, next: str = "/"):
    if not auth.configured:
        return RedirectResponse("/setup/account", status_code=303)
    if auth.valid(request.cookies.get(security.SESSION_COOKIE)):
        return RedirectResponse(security.safe_next(next), status_code=303)
    return templates.TemplateResponse(request, "login.html", {"csrf": request.state.csrf, "next": security.safe_next(next),
                                                              "error": None, "user": ""})


@app.post("/login", response_class=HTMLResponse)
def login(request: Request, username: str = Form(""), password: str = Form(""), remember: str = Form(""),
          next: str = Form("/")):
    if not auth.configured:
        return RedirectResponse("/setup/account", status_code=303)
    ip = request.client.host if request.client else "?"
    wait = auth.locked_for(ip)
    if wait:
        error = f"Too many wrong attempts. Try again in {max(1, wait // 60)} minute(s)."
    elif auth.check(ip, username.strip(), password):
        resp = RedirectResponse(security.safe_next(next), status_code=303)
        token, ttl = auth.make(long=bool(remember))
        resp.set_cookie(security.SESSION_COOKIE, token, max_age=ttl if remember else None, httponly=True,
                        samesite="lax", secure=request.url.scheme == "https", path="/")
        return resp
    else:
        left = security.MAX_FAILS - len(auth._fails.get(ip, []))
        error = "Wrong username or password." + (f" {left} attempt(s) left." if 0 < left <= 2 else "")
    return templates.TemplateResponse(request, "login.html", {"csrf": request.state.csrf, "next": security.safe_next(next),
                                                              "error": error, "user": username[:100]}, status_code=401)


@app.post("/logout")
def logout():
    resp = RedirectResponse("/login", status_code=303)
    resp.delete_cookie(security.SESSION_COOKIE, path="/")
    return resp


@app.get("/healthz")
def healthz():
    return {"ok": True}


@app.get("/", response_class=HTMLResponse)
def dashboard(request: Request):
    if not store.has_master:
        return page(request, "welcome.html", setup_started=bool((store.root / wizard.STATE).exists()))
    issues, tracks, err = [], {}, None
    try:
        master, tracks = load()
        issues = cvtracks.lint(master, tracks, base_dir=store.root)
    except (ModelError, StoreError, yaml.YAMLError) as e:
        err = str(e)
    counts = {lvl: sum(i.level == lvl for i in issues) for lvl in ("error", "warning")}
    feed_counts = feed.feed_store(store).counts()
    rows = [(t.id, t.languages, cv_i18n.text(t.title, t.languages[0])) for t in tracks.values()]
    return page(request, "dashboard.html", tracks=rows, counts=counts, load_error=err, outputs=store.outputs()[:12],
                new_matches=feed_counts.get("new", {}).get("good", 0),
                funnel=feed.feed_store(store).funnel().get("All", {"applied": 0, "interview": 0, "waiting": 0, "interview_rate": None}))


@app.post("/init")
def init_workspace():
    if store.has_master:
        return back("/", err="A master.yaml already exists")
    store.write("master.yaml", (EXAMPLES / "master.yaml").read_text(encoding="utf-8"))
    for p in (EXAMPLES / "tracks").glob("*.yaml"):
        store.write(f"tracks/{p.name}", p.read_text(encoding="utf-8"))
    return back("/", msg="Example workspace created. Edit master.yaml to make it yours.")


@app.get("/edit", response_class=HTMLResponse)
def edit(request: Request, file: str = "master.yaml"):
    try:
        text = store.read(file)
    except StoreError as e:
        return back("/", err=str(e))
    return page(request, "edit.html", file=file, text=text, files=store.files(), history=store.history(file)[:15])


def _validate(file: str, text: str) -> None:
    import tempfile
    import tomllib

    if file == "feed.toml":
        from jobfeed.config import ConfigError, parse
        try:
            parse(text)
        except ConfigError as e:
            raise ModelError(str(e)) from e
        return
    if file == "profile.toml":
        try:
            data = tomllib.loads(text)
        except tomllib.TOMLDecodeError as e:
            raise ModelError(f"TOML: {e}") from e
        if "api_key" in data.get("llm", {}):
            raise ModelError("don't store API keys here; set llm.api_key_env to the NAME of an environment variable")
        Profile.parse(data.get("countries") or ["ES"], data.get("citizenships") or [], data.get("languages"))
        return
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / Path(file).name
        p.write_text(text, encoding="utf-8")
        (load_master if file == "master.yaml" else load_track)(p)


@app.post("/edit", response_class=HTMLResponse)
def save(request: Request, file: str = Form(...), text: str = Form(...)):
    try:
        store.editable(file)
        _validate(file, text)
    except (ModelError, StoreError, ValueError, yaml.YAMLError) as e:
        return page(request, "edit.html", file=file, text=text, files=store.files(),
                    history=store.history(file)[:15], error=f"Not saved: {e}")
    store.write(file, text)
    if file == "profile.toml":
        places.ensure(store)
    return back(f"/edit?file={file}", msg="Saved (previous version kept in history)")


@app.get("/history/{name}", response_class=HTMLResponse)
def show_snapshot(request: Request, name: str):
    try:
        target, text = store.read_snapshot(name)
    except (StoreError, OSError):
        return back("/", err="No such version")
    return page(request, "snapshot.html", name=name, target=target, text=text)


@app.post("/history/{name}/restore")
def restore(name: str):
    try:
        target, text = store.read_snapshot(name)
    except (StoreError, OSError):
        return back("/", err="No such version")
    store.write(target, text)
    return back(f"/edit?file={target}", msg="Version restored (the current one was kept in history)")


@app.post("/tracks/new")
def new_track(track_id: str = Form(...), title: str = Form(...)):
    tid = track_id.strip().lower()
    if not valid_track_id(tid):
        return back("/", err="Track id: lowercase letters, digits, - and _ only")
    if store.path(f"tracks/{tid}.yaml").exists():
        return back("/", err="That track exists already")
    safe_title = title.replace('"', "'").strip() or tid
    store.write(f"tracks/{tid}.yaml", TRACK_TEMPLATE.format(id=tid, title=safe_title))
    return back(f"/edit?file=tracks/{tid}.yaml", msg="Track created")


@app.post("/tracks/{track_id}/delete")
def delete_track(track_id: str):
    try:
        store.delete_track(track_id)
    except StoreError as e:
        return back("/", err=str(e))
    return back("/", msg=f"Track {track_id} deleted (kept in history)")


@app.get("/lint", response_class=HTMLResponse)
def lint_page(request: Request):
    try:
        master, tracks = load()
        issues = cvtracks.lint(master, tracks, base_dir=store.root)
    except (ModelError, StoreError, yaml.YAMLError) as e:
        return page(request, "lint.html", issues=[], load_error=str(e))
    order = {"error": 0, "warning": 1, "info": 2}
    return page(request, "lint.html", issues=sorted(issues, key=lambda i: order[i.level]), load_error=None)


@app.post("/build")
def build(track: str = Form(""), formats: str = Form("pdf")):
    try:
        master, tracks = load()
    except (ModelError, StoreError, yaml.YAMLError) as e:
        return back("/", err=str(e))
    chosen = [track] if track else list(tracks)
    fmts = [f for f in formats.split(",") if f in RENDERERS] or ["pdf"]
    n = 0
    for tid in chosen:
        if tid not in tracks:
            return back("/", err=f"Unknown track {tid}")
        t = tracks[tid]
        opts = Options(max_pages=t.limits.get("pages"))
        for lang in t.languages:
            cv = resolve(master, t, lang)
            for fmt in fmts:
                r = RENDERERS[fmt]
                out = store.path("out") / _file_stem(t.file_name, master.person.name, tid, lang, f".{r.ext}")
                out.write_bytes(render(cv, fmt, opts))
                n += 1
    return back("/files", msg=f"Built {n} file(s)")


@app.get("/files", response_class=HTMLResponse)
def files(request: Request):
    return page(request, "files.html", outputs=store.outputs())


@app.get("/files/{rel:path}")
def get_file(rel: str, download: bool = False):
    try:
        p = store.output_path(rel)
    except StoreError:
        return Response("Not found", 404)
    media = {".pdf": "application/pdf", ".md": "text/plain; charset=utf-8", ".typ": "text/plain; charset=utf-8",
             ".json": "application/json", ".html": "text/plain; charset=utf-8"}.get(p.suffix, "application/octet-stream")
    # HTML outputs are served as text so nothing in a CV file can run as a page on this origin.
    disposition = "attachment" if download else "inline"
    return FileResponse(p, media_type=media, headers={"Content-Disposition": f'{disposition}; filename="{p.name}"'})


@app.post("/files/delete")
def delete_file(rel: str = Form(...)):
    try:
        store.output_path(rel).unlink()
    except (StoreError, OSError):
        return back("/files", err="Could not delete")
    return back("/files", msg="Deleted")


# ---------------------------------------------------------------- job check


@app.get("/check", response_class=HTMLResponse)
def check_form(request: Request, job: str = ""):
    tracks = {}
    try:
        _, tracks = load()
    except (ModelError, StoreError, yaml.YAMLError):
        pass
    form = {}
    if job:  # pre-fill from a feed item (its stored, already-sanitised text)
        j = feed.feed_store(store).job(job)
        if j:
            form = {"text": f"{j['title']}\n{j['company']}\nLocation: {j['location']}\n\n{j['text'] or ''}", "url": "",
                    "name": j["company"], "track": _track_for(j), "job": job}
    return page(request, "check.html", tracks=list(tracks), result=None, form=form)


def _tailor_cv(master, t, job_text: str, lang: str, name: str, headline: str | None = None) -> tuple[str, object, str]:
    """Tailor track `t` to a job: writes the PDF under out/tailored/. Returns (file, report, CV as Markdown).
    `headline` replaces the track title (Apply mirrors the job's own title when it's the same role)."""
    import dataclasses
    if headline:
        t = dataclasses.replace(t, title=headline)
    scores, report = cvtracks.tailor(master, t, job_text)
    cv = resolve(master, t, lang, scores)
    slug = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")[:40] or "job"
    stamp = dt.datetime.now().strftime("%Y%m%d")
    rel = f"tailored/{_file_stem(t.file_name, master.person.name, t.id, lang)}_{slug}-{stamp}.pdf"
    out = store.path(f"out/{rel}")
    out.parent.mkdir(parents=True, exist_ok=True)
    opts = Options(max_pages=t.limits.get("pages"))
    out.write_bytes(render(cv, "pdf", opts))
    return rel, report, render(cv, "md", opts).decode("utf-8")


def _track_for(job: dict) -> str:
    try:
        from jobfeed.config import parse
        cfg = parse(store.read("feed.toml"))
        return next((s.track or "" for s in cfg.searches if s.name == job["search"]), "")
    except Exception:
        return ""


@app.post("/check", response_class=HTMLResponse)
def check(request: Request, text: str = Form(""), url: str = Form(""), track: str = Form(""),
          lang: str = Form(""), use_llm: str = Form(""), name: str = Form(""), job: str = Form("")):
    job_uid = job if UID.fullmatch(job or "") else ""
    form = {"text": text, "url": url, "track": track, "lang": lang, "use_llm": use_llm, "name": name, "job": job_uid}
    tracks: dict = {}
    master = None
    try:
        master, tracks = load()
    except (ModelError, StoreError, yaml.YAMLError):
        pass
    ctx = {"tracks": list(tracks), "form": form}

    job_text = text.strip()
    if url.strip():
        try:
            job_text = fetch(url.strip())
        except FetchError as e:
            return page(request, "check.html", result=None, error=f"Couldn't fetch the URL: {e}", **ctx)
    if not job_text:
        return page(request, "check.html", result=None, error="Paste a job post or give a URL", **ctx)

    try:
        profile, llm, _ = profile_and_llm()
    except (StoreError, ValueError) as e:
        return page(request, "check.html", result=None, error=f"profile.toml: {e}", **ctx)
    job = {"description": job_text, "job_url": url.strip()}
    result = assess(job, profile, llm=llm if use_llm else None)

    tailored = None
    if track and master is not None and track in tracks:
        t = tracks[track]
        lang_ = lang if lang in t.languages else t.languages[0]
        rel, report, _ = _tailor_cv(master, t, job_text, lang_, name or result.job.get("company") or "job")
        tailored = {"report": report, "file": rel, "lang": lang_}
        if job_uid and feed.feed_store(store).job(job_uid):
            feed.feed_store(store).update(job_uid, {"cv_file": rel})
            tailored["job"] = job_uid

    return page(request, "check.html", result=result, icons=ICONS, tailored=tailored, **ctx)


# ---------------------------------------------------------------- feed

FEED_TABS = ("new", "saved", "applied", "dismissed")
UID = re.compile(r"[0-9a-f]{16}")


@app.get("/feed", response_class=HTMLResponse)
def feed_page(request: Request, status: str = "new", all: bool = False, search: str = ""):
    if status not in FEED_TABS:
        status = "new"
    fs = feed.feed_store(store)
    jobs = fs.jobs(status=status, search=search or None, good_only=not all and status == "new")
    searches = []
    try:
        from jobfeed.config import parse
        searches = [s.name for s in parse(store.read("feed.toml")).searches]
    except Exception as e:
        searches, cfg_error = [], str(e)
    else:
        cfg_error = None
    try:
        from jobfeed.config import parse as _parse
        companies = _parse(store.read("feed.toml")).companies
    except Exception:
        companies = []
    try:
        helper_cfg = _parse(store.read("feed.toml")).browser_helper
    except Exception:
        helper_cfg = {}
    return page(request, "feed.html", jobs=jobs, status=status, all=all, search=search, searches=searches, companies=companies,
                careers=fs.careers_states(), blocked=fs.blocked_hosts(), careers_labels=CAREERS_LABELS, helper=bool(helper_cfg.get("url")),
                counts=fs.counts(), last=fs.last_run(), sources=fs.source_states(), run_state=feed.state(),
                cfg_error=cfg_error, icons=ICONS, tabs=FEED_TABS)


@app.post("/feed/run")
def feed_run(force: str = Form("")):
    started = feed.run_now(store, force=bool(force))
    return back("/feed", msg="Feed run started; refresh in a minute" if started else "A run is already in progress")


@app.post("/feed/{uid}/status")
def feed_status(uid: str, status: str = Form(...), back_to: str = Form("new")):
    from jobfeed.store import STATUSES
    if status not in STATUSES or not UID.fullmatch(uid):
        return back("/feed", err="Bad request")
    feed.feed_store(store).set_status(uid, status)
    if status == "applied":
        return back(f"/jobs/{uid}", msg="Marked as applied. Add which CV you sent and any notes.")
    return back(f"/feed?status={back_to if back_to in FEED_TABS else 'new'}", msg=f"Moved to {status}")


# ---------------------------------------------------------------- tracker

def _job_or_404(uid: str):
    if not UID.fullmatch(uid):
        return None
    return feed.feed_store(store).job(uid)


@app.get("/tracker", response_class=HTMLResponse)
def tracker(request: Request, profile: str = "", status: str = ""):
    from jobfeed.store import STATUSES, TRACKED
    fs = feed.feed_store(store)
    statuses = (status,) if status in STATUSES else ("saved",) + TRACKED
    jobs = fs.tracked(statuses, profile=profile or None)
    funnel = fs.funnel()
    profiles = [p for p in funnel if p != "All"]
    return page(request, "tracker.html", jobs=jobs, funnel=funnel, profiles=profiles, profile=profile, status=status,
                statuses=STATUSES)


JOB_FIELDS = ("title", "company", "location", "url", "profile", "platform", "posted", "applied_at", "response",
              "response_at", "cv_file", "notes", "search_filter", "next_action", "next_at", "text")


def _clean_fields(form) -> dict:
    out = {}
    for k in JOB_FIELDS:
        if k in form:
            v = str(form[k]).strip()[:20000 if k == "text" else 2000]
            if k == "url" and v and not v.lower().startswith(("http://", "https://")):
                raise ValueError("Links must start with http:// or https://")
            if k in ("posted", "applied_at", "response_at", "next_at") and v:
                try:
                    dt.date.fromisoformat(v)
                except ValueError:
                    raise ValueError(f"{k}: use a date (YYYY-MM-DD)") from None
            out[k] = v
    return out


@app.get("/jobs/new", response_class=HTMLResponse)
def job_new_form(request: Request):
    return page(request, "job_new.html", profiles=[p for p in feed.feed_store(store).funnel() if p != "All"], form={})


@app.post("/jobs/new", response_class=HTMLResponse)
async def job_new(request: Request):
    form = await request.form()
    try:
        fields = _clean_fields(form)
    except ValueError as e:
        return page(request, "job_new.html", profiles=[], form=dict(form), error=str(e))
    if not (fields.get("title") or fields.get("company")):
        return page(request, "job_new.html", profiles=[], form=dict(form), error="Give at least a title or a company")
    status = str(form.get("status") or "saved")
    from jobfeed.store import STATUSES
    fields["status"] = status if status in STATUSES else "saved"
    if fields["status"] in ("applied", "interview", "offer", "rejected") and not fields.get("applied_at"):
        fields["applied_at"] = dt.date.today().isoformat()
    uid, inserted = feed.feed_store(store).add_manual_ex(fields, origin="manual")
    if inserted and fields.get("applied_at"):
        feed.feed_store(store).add_event(uid, "applied", "Added as applied")
    return back(f"/jobs/{uid}", msg="Added" if inserted else "That job is already in the tracker")


@app.get("/jobs/{uid}", response_class=HTMLResponse)
def job_page(request: Request, uid: str):
    from jobfeed.store import EVENT_KINDS, STATUSES
    j = _job_or_404(uid)
    if not j:
        return back("/tracker", err="No such job")
    return page(request, "job.html", j=j, events=feed.feed_store(store).events(uid), statuses=STATUSES,
                kinds=EVENT_KINDS, icons=ICONS, today=dt.date.today().isoformat(), has_apply=applications.get(uid) is not None)


@app.post("/jobs/{uid}")
async def job_save(request: Request, uid: str):
    if not _job_or_404(uid):
        return back("/tracker", err="No such job")
    try:
        fields = _clean_fields(await request.form())
    except ValueError as e:
        return back(f"/jobs/{uid}", err=str(e))
    feed.feed_store(store).update(uid, fields)
    return back(f"/jobs/{uid}", msg="Saved")


@app.post("/jobs/{uid}/status")
def job_status(uid: str, status: str = Form(...)):
    from jobfeed.store import STATUSES
    if status not in STATUSES or not _job_or_404(uid):
        return back("/tracker", err="Bad request")
    feed.feed_store(store).set_status(uid, status)
    return back(f"/jobs/{uid}", msg=f"Status: {status}")


@app.post("/jobs/{uid}/event")
def job_event(uid: str, kind: str = Form("note"), text: str = Form(""), at: str = Form("")):
    from jobfeed.store import EVENT_KINDS
    if not _job_or_404(uid) or kind not in EVENT_KINDS or not text.strip():
        return back(f"/jobs/{uid}", err="Write something and pick a type")
    try:
        when = dt.date.fromisoformat(at).isoformat() + "T12:00:00+00:00" if at else None
    except ValueError:
        when = None
    feed.feed_store(store).add_event(uid, kind, text.strip(), at=when)
    return back(f"/jobs/{uid}", msg="Added to the timeline")


@app.post("/jobs/{uid}/event/{event_id}/delete")
def job_event_delete(uid: str, event_id: int):
    if _job_or_404(uid):
        feed.feed_store(store).delete_event(uid, event_id)
    return back(f"/jobs/{uid}", msg="Removed")


@app.post("/jobs/{uid}/delete")
def job_delete(uid: str):
    if _job_or_404(uid):
        feed.feed_store(store).delete(uid)
        applications.delete(uid)
    return back("/tracker", msg="Deleted")


@app.get("/tracker/export.csv")
def tracker_csv():
    from jobfeed.tracker_io import export_csv
    stamp = dt.date.today().isoformat()
    return Response(export_csv(feed.feed_store(store)), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": f'attachment; filename="job-tracker-{stamp}.csv"'})


@app.get("/tracker/export.xlsx")
def tracker_xlsx():
    import io as _io

    from jobfeed.tracker_io import export_rows
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Tracker"
    for r, row in enumerate(export_rows(feed.feed_store(store)), start=1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if isinstance(value, str) and value.startswith("="):
                cell.data_type = "s"  # text, never a formula
            if r == 1:
                cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for col, width in zip("ABCDEFGHIJKLMNOPQR", (14, 14, 40, 36, 26, 20, 60, 12, 8, 14, 14, 14, 24, 11, 40, 40, 24, 14)):
        ws.column_dimensions[col].width = width
    buf = _io.BytesIO()
    wb.save(buf)
    stamp = dt.date.today().isoformat()
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="job-tracker-{stamp}.xlsx"'})


@app.post("/tracker/import")
async def tracker_import(request: Request):
    from jobfeed.tracker_io import import_csv, import_rows
    form = await request.form()
    up = form.get("file")
    if up is None or not getattr(up, "filename", ""):
        return back("/tracker", err="Choose a .csv or .xlsx file")
    data = await up.read(5_000_001)
    if len(data) > 5_000_000:
        return back("/tracker", err="File too large (5 MB max)")
    fs = feed.feed_store(store)
    try:
        if up.filename.lower().endswith(".xlsx"):
            import io as _io

            from openpyxl import load_workbook
            wb = load_workbook(_io.BytesIO(data), read_only=True, data_only=True)
            rows = list(wb.worksheets[0].iter_rows(values_only=True))
            header = [str(h or "").strip() for h in rows[0]] if rows else []
            res = import_rows(fs, [dict(zip(header, r)) for r in rows[1:]])
        else:
            res = import_csv(fs, data.decode("utf-8-sig", errors="replace"))
    except Exception as e:
        return back("/tracker", err=f"Couldn't read the file: {type(e).__name__}")
    return back("/tracker", msg=f"Imported {res['added']} job(s); {res['skipped_existing']} already there, {res['empty_rows']} empty row(s) skipped")


# ---------------------------------------------------------------- company boards

@app.post("/feed/companies/add")
def feed_company_add(url: str = Form(""), name: str = Form("")):
    from jobfeed import ats, http as fhttp
    from jobfeed.config import ConfigError, parse

    found = ats.detect(url)
    if not found:
        return back("/feed", err="Couldn't recognise that careers page. Supported: Greenhouse, Lever, Ashby, Workable, "
                                 "SmartRecruiters, Recruitee, Personio, Teamtailor, Breezy, Workday.")
    kind, slug = found
    label = (name.strip() or slug.split("/")[-1]).replace('"', "'")[:80]
    company = ats.Company(label, kind, slug)
    try:
        company.validate()
        n = len(ats.openings(company))
    except (ValueError, fhttp.FetchError) as e:
        return back("/feed", err=f"{ats.ATS_NAMES[kind]} board '{slug}' didn't answer: {e}"[:300])
    text = store.read("feed.toml")
    try:
        cfg = parse(text)
    except ConfigError as e:
        return back("/feed", err=f"feed.toml has an error, fix it first: {e}")
    if any(c.ats == kind and c.slug.lower() == slug.lower() for c in cfg.companies):
        return back("/feed", msg=f"{label} is already in your list")
    block = f'\n[[company]]\nname = "{label}"\nats = "{kind}"\nslug = "{slug}"\n'
    new = text.rstrip("\n") + "\n" + block
    if not any("companies" in s.sources for s in cfg.searches):
        note = " Add \"companies\" to a search's sources to use it."
    else:
        note = ""
    parse(new)  # must still be valid
    store.write("feed.toml", new)
    return back("/feed", msg=f"Added {label} ({ats.ATS_NAMES[kind]}, {n} open roles right now). It's included from the next run.{note}")


CAREERS_LABELS = {
    "ats": ("ok", "Uses a known job system: read through its API"),
    "jsonld": ("ok", "Job markup found: read directly"),
    "html": ("ok", "Read from the page text"),
    "browser": ("ok", "Read with the browser helper (Crawl4AI)"),
    "needs_browser": ("warn", "Needs a browser: will try the browser helper (Crawl4AI) when it's reachable"),
    "blocked": ("bad", "Refuses automated access: skipped (never bypassed)"),
    "robots": ("bad", "robots.txt says no: skipped"),
    "error": ("bad", "Error"),
    "empty": ("warn", "No jobs found"),
}


@app.post("/feed/careers/add")
def feed_careers_add(url: str = Form(""), name: str = Form("")):
    from jobfeed.config import ConfigError, parse

    url = url.strip()
    if not url.lower().startswith(("http://", "https://")):
        return back("/feed", err="Give the company website or careers page, starting with https://")
    from urllib.parse import urlparse
    label = (name.strip() or urlparse(url).netloc.removeprefix("www.")).replace('"', "'")[:80]
    text = store.read("feed.toml")
    try:
        cfg = parse(text)
    except ConfigError as e:
        return back("/feed", err=f"feed.toml has an error, fix it first: {e}")
    if any(c["url"].rstrip("/") == url.rstrip("/") for c in cfg.careers):
        return back("/feed", msg=f"{label} is already in your list")
    safe_url = url.replace('"', "%22").replace("\\", "%5C")
    new = text.rstrip("\n") + f'\n\n[[careers]]\nname = "{label}"\nurl = "{safe_url}"\n'
    try:
        parse(new)
    except ConfigError as e:
        return back("/feed", err=str(e))
    store.write("feed.toml", new)
    note = "" if any("careers" in s.sources for s in cfg.searches) else ' Add "careers" to a search\'s sources to use it.'
    return back("/feed", msg=f"Added {label}. It's checked on the next run (Run now with \"ignore source intervals\" to check it now).{note}")


# ---------------------------------------------------------------- apply


def _apply_context(j: dict, track: str = "", lang: str = ""):
    """Load everything Apply needs for a job: (master, raw master dict, profile, track, lang)."""
    master, tracks = load()
    if not tracks:
        raise StoreError("no CV tracks yet")
    tid = track if track in tracks else (_track_for(j) if _track_for(j) in tracks else next(iter(tracks)))
    t = tracks[tid]
    lang = lang if lang in t.languages else apply.guess_lang(j.get("text") or "", t.languages)
    return master, yaml.safe_load(store.read("master.yaml")), store.profile(), t, lang


def _prepare(uid: str, j: dict, a: "apply.Application", qs: list, extra: list[str], track: str, lang: str) -> str | None:
    """Tailor the CV, fill by code, start drafting. Returns an error message for the page, or None."""
    try:
        master, raw, prof, t, lang = _apply_context(j, track, lang)
    except (ModelError, StoreError, yaml.YAMLError) as e:
        return f"CV data: {e}"
    rel, _, cv_md = _tailor_cv(master, t, j.get("text") or j.get("title") or "", lang, j.get("company") or "job",
                               headline=apply.headline_for(j.get("title") or "", t.id))
    feed.feed_store(store).update(uid, {"cv_file": rel})
    old = next((q for q in a.questions if q.id == apply.COVER_LETTER_ID), None)
    qs = [q for q in qs if q.id != apply.COVER_LETTER_ID]
    a.track, a.lang, a.cv_file, a.questions = t.id, lang, rel, qs
    a.ai_ban = apply.ai_ban(extra + [q.label for q in qs] + [j.get("text") or ""])
    cover = old if old and old.origin == "edited" else apply.cover_letter_question(lang)
    if a.ai_ban and cover.origin != "edited":
        cover.origin, cover.note = "you", "This employer forbids AI-written text: write it yourself, or skip it."
    a.questions.append(cover)
    a.status, a.error = "ready", ""
    apply.fill_by_code(a.questions, apply.facts(raw, prof, j, lang, rel, t.id))
    applications.save(a)
    if a.ai_ban or not any(q.origin == "ai" for q in a.questions):
        return None
    llm = apply.LLM.from_profile(prof, llm_key())
    if llm is None:
        a.status, a.error = "error", "No local model configured: set [llm] base_url (or [apply] ollama_url) in profile.toml."
        applications.save(a)
        return None
    apply.draft_async(applications, uid, llm, _with_about(cv_md), _job_line(j))
    return None


def _with_about(cv_md: str) -> str:
    """The CV plus the candidate's own notes (data/about.md, synced from the second brain), if present."""
    try:
        about = store.path("about.md").read_text(encoding="utf-8").strip()
    except (OSError, StoreError):
        about = ""
    about = re.sub(r"^---\n.*?\n---\n", "", about, flags=re.S)[:12000]  # drop Obsidian frontmatter
    return f"{cv_md}\n\n# More facts from the candidate's own notes\n\n{about}" if about else cv_md


def _job_line(j: dict) -> str:
    text = re.sub(r"\s+", " ", j.get("text") or "")[:2500]
    return f"{j.get('title') or ''} at {j.get('company') or ''}. {text}".strip()


@app.post("/jobs/{uid}/apply")
def apply_start(uid: str, track: str = Form(""), lang: str = Form("")):
    j = _job_or_404(uid)
    if not j:
        return back("/feed", err="No such job")
    if apply.is_running(uid):
        return back(f"/jobs/{uid}/apply", msg="Still drafting; refresh in a moment")
    a = applications.get(uid) or apply.Application(uid=uid)
    note = None
    try:
        source, form_url, qs, extra = apply.read_form(j.get("url") or "")
    except FetchError as e:
        source, form_url, qs, extra = "none", "", [], []
        note = f"Couldn't read the application form ({e}). Paste the questions below."
    if source == "none" and a.source == "pasted":  # keep questions pasted earlier
        source, qs, extra = "pasted", a.questions, []
    a.source, a.form_url = source, form_url
    err = _prepare(uid, j, a, qs, extra, track, lang)
    if err:
        return back(f"/jobs/{uid}", err=err)
    if source == "none" and not note:
        note = "This job's form can't be read automatically. Paste its questions below."
    return back(f"/jobs/{uid}/apply", msg=note or "Form read, CV tailored. Open answers are being drafted.")


@app.post("/jobs/{uid}/apply/questions")
def apply_questions(uid: str, questions: str = Form("")):
    j = _job_or_404(uid)
    a = applications.get(uid) if j else None
    if not a:
        return back(f"/jobs/{uid}", err="Start with Apply first")
    if apply.is_running(uid):
        return back(f"/jobs/{uid}/apply", err="Wait until the current drafting finishes")
    qs = apply.parse_pasted(questions[:20000])
    if not qs:
        return back(f"/jobs/{uid}/apply", err="Paste at least one question")
    a.source = "pasted"
    err = _prepare(uid, j, a, qs, [], a.track, a.lang)
    return back(f"/jobs/{uid}/apply", err=err, msg=None if err else f"{len(qs)} question(s) added")


@app.get("/jobs/{uid}/apply", response_class=HTMLResponse)
def apply_page(request: Request, uid: str):
    j = _job_or_404(uid)
    a = applications.get(uid) if j else None
    if not a:
        return back(f"/jobs/{uid}" if j else "/feed", err="Press Apply first")
    if a.status == "drafting" and not apply.is_running(uid):
        a.status, a.error = "error", "Drafting was interrupted (the app restarted). Draft again."
        applications.save(a)
    try:
        _, tracks = load()
    except (ModelError, StoreError, yaml.YAMLError):
        tracks = {}
    langs = tracks[a.track].languages if a.track in tracks else [a.lang]
    return page(request, "apply.html", j=j, a=a, tracks=list(tracks), langs=langs,
                counts={k: sum(q.origin == k for q in a.questions) for k in ("code", "ai", "you", "edited")},
                needs=sum("[NEEDS INPUT" in q.answer for q in a.questions))


@app.post("/jobs/{uid}/apply/save")
async def apply_save(request: Request, uid: str):
    a = applications.get(uid) if _job_or_404(uid) else None
    if not a:
        return back("/feed", err="No such application")
    form = await request.form()
    for i, q in enumerate(a.questions):
        key = f"a{i}"
        if key in form:
            v = str(form[key]).strip()[:8000]
            if v != q.answer:
                q.answer, q.origin = v, "edited"
    applications.save(a)
    return back(f"/jobs/{uid}/apply", msg="Answers saved")


@app.post("/jobs/{uid}/apply/draft")
def apply_redraft(uid: str, index: int = Form(-1)):
    j = _job_or_404(uid)
    a = applications.get(uid) if j else None
    if not a:
        return back("/feed", err="No such application")
    if a.ai_ban:
        return back(f"/jobs/{uid}/apply", err="This employer forbids AI-written answers, so nothing is drafted.")
    only = index if 0 <= index < len(a.questions) else None
    if only is not None:
        a.questions[only].origin = "ai"
    else:
        for q in a.questions:
            if q.origin == "ai":
                q.answer = ""
    applications.save(a)
    try:
        master, _, prof, t, lang = _apply_context(j, a.track, a.lang)
    except (ModelError, StoreError, yaml.YAMLError) as e:
        return back(f"/jobs/{uid}/apply", err=f"CV data: {e}")
    llm = apply.LLM.from_profile(prof, llm_key())
    if llm is None:
        return back(f"/jobs/{uid}/apply", err="No local model configured in profile.toml")
    cv_md = render(resolve(master, t, lang, cvtracks.tailor(master, t, j.get("text") or "")[0]), "md").decode("utf-8")
    if not apply.draft_async(applications, uid, llm, _with_about(cv_md), _job_line(j), only):
        return back(f"/jobs/{uid}/apply", err="Already drafting")
    return back(f"/jobs/{uid}/apply", msg="Drafting…")


from . import wizard  # noqa: E402  (first-run account and setup wizard)

app.include_router(wizard.router)
